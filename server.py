from flask import Flask, request, send_file, render_template
from pptx import Presentation
from pptx.util import Inches, Pt # Import positioning utilities
from io import BytesIO
import requests # --- NEW: Import requests library for fetching URLs ---
pip install Flask python-pptx requests

app = Flask(__name__)

# --- NEW: Helper function to fetch image from URL into memory ---
def fetch_image_into_memory(url):
    try:
        # Important: Change 'blob' to 'raw' in GitHub URLs to get the actual image data
        if "github.com" in url and "/blob/" in url:
             url = url.replace("/blob/", "/raw/")

        print(f"Attempting to fetch image from: {url}")
        response = requests.get(url, timeout=10) # Set a timeout so it doesn't hang forever
        response.raise_for_status() # Raise an exception if the download failed (e.g., 404 error)

        # Wrap the raw binary data in a BytesIO stream so pptx can read it
        image_stream = BytesIO(response.content)
        return image_stream
    except requests.exceptions.RequestException as e:
        print(f"Error downloading image: {e}")
        return None
# ----------------------------------------------------------------

@app.route('/')
def index():
    # Ensure your HTML file is saved as 'templates/index.html'
    return render_template('index.html')

@app.route('/upload-and-design', methods=['POST'])
def upload_and_design():
    if 'pptFile' not in request.files:
        return "No file part", 400

    file = request.files['pptFile']
    if file.filename == '':
        return "No selected file", 400

    if file:
        try:
            # 1. Read the incoming PPTX file into memory
            prs = Presentation(file)

            # --- NEW: Fetch the image ONCE before iterating slides ---
            # This prevents downloading it repeatedly for every slide.
            image_url = "https://github.com/anusin1805/DesignPPTX/blob/main/gender%20graph%20difference.png"
            image_memory_stream = fetch_image_into_memory(image_url)

            if not image_memory_stream:
                print("Warning: Could not fetch image. Slides will be generated without it.")


            # 2. Iterate through slides and apply designs
            for slide_idx, slide in enumerate(prs.slides):

                # --- Existing Text Design (keeping this as an example) ---
                # Place text at top left
                left_text = Inches(0.5)
                top_text = Inches(0.5)
                width_text = Inches(8)
                height_text = Inches(1)
                textbox = slide.shapes.add_textbox(left_text, top_text, width_text, height_text)
                p = textbox.text_frame.add_paragraph()
                p.text = f"Design Trace added to Slide {slide_idx + 1}"
                p.font.size = Pt(18)
                p.font.color.rgb = (0x55, 0x55, 0x55) # Dark gray

                # --- NEW: IMAGE ADDITION LOGIC ---
                if image_memory_stream:
                    # CRITICAL: Reset stream pointer to the beginning for each slide read
                    image_memory_stream.seek(0)

                    # Define image position and size
                    # Let's place it below the text, centered horizontally
                    img_left = Inches(2.5)   # 2.5 inches from left edge
                    img_top = Inches(2.0)    # 2.0 inches from top edge
                    img_width = Inches(5.0)  # 5 inches wide
                    # We omit height so pptx calculates it automatically based on aspect ratio

                    try:
                        slide.shapes.add_picture(
                            image_memory_stream,
                            img_left,
                            img_top,
                            width=img_width
                            # height=img_height # Optional
                        )
                    except Exception as e:
                        print(f"Error adding picture to slide {slide_idx}: {e}")
                # -------------------------------------


            # 3. Save the modified presentation to memory
            output_pptx = BytesIO()
            prs.save(output_pptx)
            output_pptx.seek(0) # Rewind to start of stream before sending

            # 4. Send back to client
            return send_file(
                output_pptx,
                mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation',
                as_attachment=True,
                download_name='designed_with_image.pptx'
            )

        except Exception as e:
            # Print error to console for debugging and send to user
            print(f"Server Error: {str(e)}")
            return f"Error processing file: {str(e)}", 500

if __name__ == '__main__':
    # Make sure you have a folder named 'templates' next to this script
    # and put your HTML file inside it as 'index.html'
    app.run(debug=True, port=5000)
